"""Legacy BIFF cell reader; reuse the XLSX grid and provenance representation."""
import xlrd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

from .xlsx import convert_workbook


def convert_xls(path, client=None, model=None, prompt=None, workers=4):
    source = xlrd.open_workbook(str(path), formatting_info=True)
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904 if source.datemode else CALENDAR_WINDOWS_1900
    book.remove(book.active)
    stored = {}
    try:
        for original in source.sheets():
            sheet = book.create_sheet(original.name)
            sheet.sheet_state = {0: 'visible', 1: 'hidden', 2: 'veryHidden'}.get(original.visibility, 'hidden')
            raw_cells = {}
            stored[original.name] = raw_cells
            for row in range(original.nrows):
                for column in range(original.row_len(row)):
                    raw = original.cell(row, column)
                    if raw.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    cell = sheet.cell(row + 1, column + 1)
                    value = raw.value
                    if raw.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, source.datemode)
                    elif raw.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif raw.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, f'#UNKNOWN_ERROR_{value}')
                    elif raw.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                        value = int(value)
                    cell.value = value
                    # A text cell beginning '=' is text, never an executable formula.
                    if raw.ctype == xlrd.XL_CELL_TEXT:
                        cell.data_type = 's'
                    elif raw.ctype == xlrd.XL_CELL_ERROR:
                        cell.data_type = 'e'
                    xf = source.xf_list[raw.xf_index]
                    cell.number_format = source.format_map[xf.format_key].format_str
                    raw_cells[cell.coordinate] = dict(value=raw.value, xml_type=f'xls_type_{raw.ctype}',
                                                       style_index=raw.xf_index, raw_value=raw.value)
            for (row, column), note in original.cell_note_map.items():
                sheet.cell(row + 1, column + 1).comment = Comment(note.text, note.author or '')
            for (row, column), link in original.hyperlink_map.items():
                target = link.url_or_path or link.textmark
                if target:
                    cell = sheet.cell(row + 1, column + 1)
                    original_value = cell.value
                    cell.hyperlink = target
                    cell.value = original_value
            for row, info in original.rowinfo_map.items():
                sheet.row_dimensions[row + 1].hidden = bool(info.hidden)
            for column, info in original.colinfo_map.items():
                sheet.column_dimensions[get_column_letter(column + 1)].hidden = bool(info.hidden)
            # Capture raw values first: merging intentionally discards covered display values.
            for r0, r1, c0, c1 in original.merged_cells:
                sheet.merge_cells(start_row=r0 + 1, end_row=r1,
                                  start_column=c0 + 1, end_column=c1)
        text, document = convert_workbook(
            book, None, stored, client, model, prompt, workers,
            extra_warnings=[
                '.xls 返回单元格保存值；公式只能读取缓存结果，不能可靠区分公式单元格或恢复公式表达式，未重新计算。',
                '.xls 的嵌入图片、图表和形状暂未提取或 OCR；若材料含此类内容，必须核对原件。',
            ],
        )
        document['format'] = 'xls'
        document['date_mode'] = source.datemode
        document['formula_support'] = 'cached_values_only'
        document['media_support'] = 'not_extracted'
        return text, document
    finally:
        source.release_resources()
        book.close()

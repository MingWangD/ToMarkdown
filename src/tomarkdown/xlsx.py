"""Read spreadsheet cells directly; never infer headers or fill unmerged blanks."""
import hashlib
import html
import io
import re
import posixpath
from concurrent.futures import ThreadPoolExecutor
from contextlib import nullcontext
from datetime import date, datetime, time, timedelta
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from markitdown_ocr import LLMVisionOCRService

from .ocr import ordered_window


def _stored_cells(path):
    """OOXML may contain values covered by merges; openpyxl intentionally discards them."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    result = {}
    with ZipFile(path) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            strings = [''.join(t.text or '' for t in si.findall('.//s:t', ns))
                       for si in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
        relationships = {r.attrib['Id']: r.attrib['Target'] for r in
                         ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        for sheet in workbook.findall('s:sheets/s:sheet', ns):
            target = relationships[sheet.attrib[rel_ns]]
            part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            values = {}
            for c in ET.fromstring(archive.read(part)).findall('s:sheetData/s:row/s:c', ns):
                v, f = c.find('s:v', ns), c.find('s:f', ns)
                kind = c.attrib.get('t', 'n')
                if kind == 's' and v is not None:
                    value = strings[int(v.text)]
                elif kind == 'inlineStr':
                    value = ''.join(t.text or '' for t in c.findall('s:is//s:t', ns))
                elif v is not None and v.text is not None:
                    value = v.text
                    if kind == 'b':
                        value = value == '1'
                    elif kind == 'n':
                        value = int(value) if re.fullmatch(r'[+-]?\d+', value) else float(value)
                elif f is not None:
                    value = None
                else:
                    continue
                values[c.attrib['r']] = dict(value=value, xml_type=kind,
                                             style_index=int(c.attrib.get('s', '0')),
                                             raw_value=v.text if v is not None else None,
                                             formula_text=f.text if f is not None else None,
                                             formula_attributes=dict(f.attrib) if f is not None else None)
            result[sheet.attrib['name']] = values
    return result


def json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"total_seconds": value.total_seconds()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def escape_cell(value):
    text = html.escape(str(value), quote=False).replace('\\', '\\\\')
    text = re.sub(r'([|`*_\[\]])', r'\\\1', text)
    return text.replace('\r\n', '\n').replace('\r', '\n').replace('\n', '<br>')


def display_value(cell):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type == 'f':
        return f'公式：{value}（结果见来源记录；未重新计算）'
    # Preserve common numeric identifiers without pretending to implement Excel's renderer.
    if (isinstance(value, int) and not isinstance(value, bool)
            and value >= 0 and re.fullmatch(r'0+', cell.number_format)):
        return str(value).zfill(len(cell.number_format))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def convert_xlsx(path, client=None, model=None, prompt=None, workers=4):
    stored = _stored_cells(path)
    book = load_workbook(path, data_only=False)
    try:
        cached = load_workbook(path, data_only=True)
    except BaseException:
        book.close()
        raise
    return convert_workbook(book, cached, stored, client, model, prompt, workers,
                            media_archive=path)


def convert_workbook(book, cached, stored, client=None, model=None, prompt=None,
                     workers=4, *, media_archive=None, extra_warnings=()):
    """Shared faithful grid renderer for native XLSX and normalized legacy XLS cells."""
    warnings = [
        '保留原值及数字格式，不保证复现全部 Excel 显示格式；公式缓存未经重新计算。',
        '图表、形状、印章含义和图片实体计数未核验；解析成功不代表材料合格。',
        *extra_warnings,
    ]
    output = ['> Excel 原始网格：首列为行号，其余列为 Excel 列号；表头保留为原始行。',
              '> ↳ 表示引用合并区域左上角单元格，不是新增记录；空白保持空白。']
    sheets = []
    try:
        for sheet in book.worksheets:
            # Iterating stored cells avoids allocating a huge empty range due to formatting.
            native = sorted((c for c in sheet._cells.values()
                             if c.value is not None or c.comment or c.hyperlink),
                            key=lambda c: (c.row, c.column))
            merges = sorted(sheet.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))
            rows = {cell.row for cell in native}
            columns = {cell.column for cell in native}
            for merged in merges:
                rows.update((merged.min_row, merged.max_row))
                columns.update((merged.min_col, merged.max_col))
            # Retain modest empty gaps; unusually sparse ranges keep explicit coordinates.
            if rows and max(rows) - min(rows) < 1000:
                rows = set(range(min(rows), max(rows) + 1))
            if columns and max(columns) - min(columns) < 256:
                columns = set(range(min(columns), max(columns) + 1))
            row_numbers, column_numbers = sorted(rows), sorted(columns)
            records = []
            hidden_rows = sorted(r for r, dim in sheet.row_dimensions.items() if dim.hidden)
            hidden_columns = [name for name, dim in sheet.column_dimensions.items() if dim.hidden]
            sheet_data = dict(name=sheet.title, state=sheet.sheet_state,
                              merged_ranges=[str(r) for r in merges], cells=records,
                              hidden_rows=hidden_rows, hidden_columns=hidden_columns,
                              rendered_rows=row_numbers,
                              rendered_columns=[get_column_letter(c) for c in column_numbers])
            sheets.append(sheet_data)
            if sheet.sheet_state != 'visible' or hidden_rows or hidden_columns:
                warnings.append(f'工作表 {sheet.title} 含隐藏内容；已保留单元格值及隐藏状态。')
            by_position = {(cell.row, cell.column): cell for cell in native}
            for cell in native:
                record = dict(coordinate=cell.coordinate, value=json_value(cell.value),
                              data_type=cell.data_type, number_format=cell.number_format)
                if cell.data_type == 'f':
                    record['formula'] = json_value(cell.value)
                    record['cached_value'] = json_value(cached[sheet.title][cell.coordinate].value)
                    warnings.append(f'{sheet.title}!{cell.coordinate}：公式缓存未验证，可能缺失或过期。')
                for merged in merges:
                    if cell.coordinate in merged:
                        record['merged_range'] = str(merged)
                        break
                if cell.hyperlink:
                    record['hyperlink'] = cell.hyperlink.target or cell.hyperlink.location
                if cell.comment:
                    record['comment'] = cell.comment.text
                records.append(record)
            covered = []
            visible_coordinates = {cell.coordinate for cell in native}
            for coordinate, raw in stored.get(sheet.title, {}).items():
                if coordinate in visible_coordinates:
                    continue
                for merged in merges:
                    if coordinate in merged:
                        anchor = sheet.cell(merged.min_row, merged.min_col)
                        record = dict(coordinate=coordinate, value=raw['value'],
                                      data_type=raw['xml_type'], raw_storage=raw,
                                      visibility='covered_by_merge', merged_range=str(merged),
                                      merged_anchor=anchor.coordinate)
                        records.append(record)
                        covered.append(record)
                        break
            if covered:
                warnings.append(f'{sheet.title}：有 {len(covered)} 个底层存储值被合并区域遮盖，已单独保留；不得当作正常显示值。')
                if any(c['value'] != sheet[c['merged_anchor']].value
                       for c in covered if sheet[c['merged_anchor']].value is not None):
                    warnings.append(f'{sheet.title}：被合并遮盖的值与锚点值存在差异，需要核对原件和业务含义。')
            output.append(f'## {escape_cell(sheet.title)}')
            output.append(f'工作表状态：{sheet.sheet_state}。合并范围：'
                          + (', '.join(map(str, merges)) or '无') + '。')
            output.append('')
            if covered:
                output.extend(['### 被合并区域遮盖的底层值（需复核）', '',
                               '| 原坐标 | 合并锚点 | 原存储值 |', '| --- | --- | --- |'])
                output.extend(f'| {c["coordinate"]} | {c["merged_anchor"]} | {escape_cell(c["value"])} |'
                              for c in covered)
                output.append('')
            if not row_numbers or not column_numbers:
                output.append('（空工作表）')
                continue
            output.extend(['### 显示网格', ''])
            output.extend([
                '| 行 | ' + ' | '.join(get_column_letter(c) for c in column_numbers) + ' |',
                '| --- | ' + ' | '.join('---' for _ in column_numbers) + ' |',
            ])
            for row in row_numbers:
                values = []
                for column in column_numbers:
                    cell = by_position.get((row, column))
                    if cell is not None:
                        value = escape_cell(display_value(cell))
                    else:
                        value = ''
                        for merged in merges:
                            if (merged.min_row <= row <= merged.max_row
                                    and merged.min_col <= column <= merged.max_col):
                                anchor = sheet.cell(merged.min_row, merged.min_col)
                                value = f'↳ {anchor.coordinate}'
                                # References avoid duplicating long text as apparent extra records.
                                break
                    values.append(value)
                output.append(f'| {row} | ' + ' | '.join(values) + ' |')
            output.append('')
        images = []
        with ZipFile(media_archive) if media_archive is not None else nullcontext() as archive:
            parts = sorted(name for name in archive.namelist()
                           if name.startswith('xl/media/') and not name.endswith('/')) if archive else []
            def recognize(part):
                data = archive.read(part)
                item = dict(part=part, sha256=hashlib.sha256(data).hexdigest(),
                            size_bytes=len(data), location_granularity='archive_part')
                if client is None:
                    item.update(status='not_ocr', text=None)
                elif part.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp')):
                    service = LLMVisionOCRService(client=client, model=model, default_prompt=prompt)
                    result = service.extract_text(io.BytesIO(data))
                    if result.error:
                        raise RuntimeError(f'{part} OCR 识别失败：{result.error}')
                    item.update(status='ocr_returned' if result.text.strip() else 'ocr_empty', text=result.text)
                else:
                    item.update(status='unsupported_media', text=None)
                return item
            with ThreadPoolExecutor(max_workers=workers) as executor:
                images = list(ordered_window(executor, recognize, parts, workers))
        if images:
            warnings.append('嵌入媒体按原始文件部分保留身份；尚未映射到工作表锚点，不能据此统计业务图片数量。')
            for image in images:
                output.append(f'### 嵌入媒体 {escape_cell(image["part"])}')
                output.append(f'处理状态：{image["status"]}。')
                if image['text']:
                    output.append(image['text'])
                if image['status'] != 'ocr_returned':
                    warnings.append(f'{image["part"]}：{image["status"]}，需要人工检查原图。')
        document = dict(sheets=sheets, images=images, warnings=warnings,
                        date_epoch=book.epoch.isoformat())
        output.extend(['## 转换范围与待复核项', *('- ' + escape_cell(w) for w in warnings)])
        return '\n\n'.join(output[:2]) + '\n\n' + '\n'.join(output[2:]), document
    finally:
        book.close()
        if cached is not None:
            cached.close()

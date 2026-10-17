"""Evidence fidelity through the public conversion API, without a model service."""
import hashlib
import base64
import io
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook

from tomarkdown import convert
from test_ocr_performance import VisionEndpoint


class DocumentEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '台账'
        sheet.append(['企业', '项目', '津贴', '备注'])
        sheet.append(['甲公司', 'JAVA | 鼎瑞\n第二行', 7000, 'NaN'])
        sheet.append([None, '第二岗位', 0, None])
        sheet.merge_cells('A2:A3')
        sheet['B4'] = '真正空白的企业单元格在 A4'
        path = self.root / 'ledger.xlsx'
        book.save(path)
        return path

    def test_grid_escapes_pipes_and_distinguishes_blank_from_merge(self):
        result = convert(self.workbook(), ocr=False, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        text = result.files[0].read_text()
        self.assertIn('JAVA \\| 鼎瑞<br>第二行', text)
        self.assertIn('↳ A2', text)
        self.assertNotIn('Unnamed:', text)
        rows = [re.split(r'(?<!\\)\|', line)[1:-1]
                for line in text.splitlines() if line.startswith('|')]
        self.assertEqual({len(row) for row in rows}, {5})
        fourth = next(row for row in rows if row[0].strip() == '4')
        self.assertEqual(fourth[1].strip(), '')

    def test_evidence_preserves_cells_merges_types_and_source_identity(self):
        source = self.workbook()
        result = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        evidence = json.loads(result.evidence_files[source].read_text())
        self.assertEqual(evidence['source']['sha256'], hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertEqual(evidence['output']['sha256'], hashlib.sha256(result.files[0].read_bytes()).hexdigest())
        sheet = evidence['document']['sheets'][0]
        self.assertEqual(sheet['merged_ranges'], ['A2:A3'])
        cells = {cell['coordinate']: cell for cell in sheet['cells']}
        self.assertEqual(cells['B2']['value'], 'JAVA | 鼎瑞\n第二行')
        self.assertEqual(cells['D2']['value'], 'NaN')
        self.assertEqual(cells['C3']['value'], 0)
        self.assertEqual(cells['A2']['merged_range'], 'A2:A3')
        self.assertNotIn('A3', cells)  # Missing native cells are not invented facts.
        self.assertEqual(evidence['location_granularity'], 'cell')

    def test_formulas_hidden_sheets_formats_and_comments_are_explicit(self):
        book = Workbook()
        sheet = book.active
        sheet['A1'] = 123
        sheet['A1'].number_format = '000000'
        sheet['B1'] = '=SUM(A1,1)'
        sheet.row_dimensions[1].hidden = True
        book.create_sheet('内部').sheet_state = 'hidden'
        path = self.root/'formula.xlsx'
        book.save(path)
        result = convert(path, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        data = json.loads(result.evidence_files[path].read_text())
        sheets = data['document']['sheets']
        self.assertEqual(len(sheets), 2)
        self.assertEqual(sheets[1]['state'], 'hidden')
        cells = {cell['coordinate']: cell for cell in sheets[0]['cells']}
        self.assertEqual(cells['A1']['number_format'], '000000')
        self.assertEqual(cells['B1']['formula'], '=SUM(A1,1)')
        self.assertIsNone(cells['B1']['cached_value'])
        self.assertEqual(sheets[0]['hidden_rows'], [1])
        self.assertTrue(any('缓存' in warning for warning in data['warnings']))
        self.assertIn('000123', result.files[0].read_text())

    def test_evidence_mode_does_not_overwrite_same_stem_materials(self):
        folder = self.root/'materials'
        folder.mkdir()
        (folder/'same.txt').write_text('text evidence')
        (folder/'same.html').write_text('<p>html evidence</p>')
        result = convert(folder, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        self.assertEqual(len(set(result.files)), 2)
        self.assertEqual(len(set(result.evidence_files.values())), 2)
        self.assertEqual({path.read_text().strip() for path in result.files},
                         {'text evidence', 'html evidence'})
        for source, path in result.evidence_files.items():
            data = json.loads(path.read_text())
            self.assertEqual(data['source']['name'], source.name)
            self.assertEqual(data['location_granularity'], 'file')

    def test_evidence_failure_is_not_published_as_success(self):
        path = self.root/'broken.xlsx'
        path.write_text('not an xlsx')
        result = convert(path, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.files)
        self.assertFalse(result.evidence_files)
        self.assertIn(path, result.errors)
        self.assertFalse(list((self.root/'out').glob('*.json')))

    def test_embedded_image_original_bytes_and_failure_are_preserved(self):
        from PIL import Image
        from openpyxl.drawing.image import Image as SheetImage
        image = self.root/'stamp.png'
        Image.new('RGB', (1024, 640), 'white').save(image)
        book = Workbook()
        book.active['A1'] = '图片是材料的一部分'
        book.active.add_image(SheetImage(io.BytesIO(image.read_bytes())), 'B3')
        path = self.root/'image.xlsx'
        book.save(path)
        for fail in (False, True):
            endpoint = VisionEndpoint(delay=0, fail=fail)
            with self.subTest(fail=fail), patch('tomarkdown.core.OpenAI', return_value=endpoint):
                result = convert(path, api_key='fake', evidence=True,
                                 results_dir=self.root/str(fail))
                self.assertEqual(bool(result.errors), fail)
                if fail:
                    self.assertFalse(result.evidence_files)
                    continue
                payload = endpoint.payloads[0]['messages'][0]['content'][1]['image_url']['url']
                self.assertEqual(base64.b64decode(payload.split(',', 1)[1]), image.read_bytes())
                data = json.loads(result.evidence_files[path].read_text())
                self.assertEqual(data['document']['images'][0]['status'], 'ocr_returned')
                self.assertEqual(data['document']['images'][0]['location_granularity'], 'archive_part')

    def test_source_change_during_conversion_is_rejected(self):
        source = self.workbook()
        from tomarkdown.xlsx import convert_xlsx
        def changed(*args, **kwargs):
            converted = convert_xlsx(*args, **kwargs)
            source.write_bytes(source.read_bytes() + b'changed')
            return converted
        with patch('tomarkdown.core.convert_xlsx', side_effect=changed):
            result = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.files)
        self.assertFalse(result.evidence_files)
        self.assertIn('源文件发生变化', result.errors[source])

    def test_stored_values_covered_by_merges_are_not_silently_discarded(self):
        source = self.workbook()
        with ZipFile(source) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        root = ET.fromstring(contents['xl/worksheets/sheet1.xml'])
        row = root.find(f'{{{ns}}}sheetData/{{{ns}}}row[@r="3"]')
        cell = ET.SubElement(row, f'{{{ns}}}c', {'r': 'A3', 't': 'inlineStr'})
        inline = ET.SubElement(cell, f'{{{ns}}}is')
        ET.SubElement(inline, f'{{{ns}}}t').text = '被合并遮盖的不同企业名称'
        contents['xl/worksheets/sheet1.xml'] = ET.tostring(root)
        with ZipFile(source, 'w') as archive:
            for name, data in contents.items():
                archive.writestr(name, data)
        result = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        data = json.loads(result.evidence_files[source].read_text())
        cells = {c['coordinate']: c for c in data['document']['sheets'][0]['cells']}
        self.assertEqual(cells['A3']['value'], '被合并遮盖的不同企业名称')
        self.assertEqual(cells['A3']['merged_anchor'], 'A2')
        self.assertEqual(cells['A3']['visibility'], 'covered_by_merge')
        self.assertIn('被合并遮盖的不同企业名称', result.files[0].read_text())
        self.assertTrue(any('差异' in w for w in data['warnings']))

    def test_blank_cell_annotations_are_retained(self):
        from openpyxl.comments import Comment
        book = Workbook()
        book.active['A1'].comment = Comment('必须补交合同', 'Reviewer')
        source = self.root/'comment.xlsx'
        book.save(source)
        result = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertFalse(result.errors)
        data = json.loads(result.evidence_files[source].read_text())
        self.assertEqual(data['document']['sheets'][0]['cells'][0]['comment'], '必须补交合同')

    def test_manifest_publication_failure_never_leaves_mismatched_pair(self):
        import os
        source = self.workbook()
        first = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        manifest = first.evidence_files[source]
        output = first.files[0]
        original_replace = os.replace
        def fail_manifest(src, dst):
            if Path(dst) == manifest:
                raise OSError('simulated publication failure')
            original_replace(src, dst)
        with patch('tomarkdown.evidence.os.replace', side_effect=fail_manifest):
            result = convert(source, ocr=False, evidence=True, results_dir=self.root/'out')
        self.assertIn('simulated publication failure', result.errors[source])
        self.assertFalse(result.files)
        self.assertFalse(manifest.exists())
        self.assertFalse(output.exists())
        self.assertFalse(list(output.parent.glob('.tomarkdown-*')))


if __name__ == '__main__':
    unittest.main()

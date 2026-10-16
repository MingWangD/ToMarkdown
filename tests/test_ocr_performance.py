"""Exercise real converters with a deterministic, delayed vision endpoint."""
import hashlib
import base64
import io
import tempfile
import threading
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image

from tomarkdown import convert
from tomarkdown.ocr import OCRClient
from tomarkdown.core import OCR_PROMPT
from markitdown import MarkItDown
import pdfplumber


class VisionEndpoint:
    def __init__(self, delay=0.04, fail=False):
        self.delay, self.fail = delay, fail
        self.calls = self.active = self.peak = 0
        self.lock = threading.Lock()
        self.payloads = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    def __enter__(self):
        return self

    def __exit__(self, *args):
        pass

    def create(self, **kwargs):
        with self.lock:
            self.calls += 1
            self.payloads.append(kwargs)
            self.active += 1
            self.peak = max(self.peak, self.active)
        try:
            time.sleep(self.delay)
            if self.fail:
                raise RuntimeError('synthetic OCR failure')
            data = kwargs['messages'][0]['content'][1]['image_url']['url']
            content = 'OCR-' + hashlib.sha256(data.encode()).hexdigest()[:12]
            return SimpleNamespace(choices=[SimpleNamespace(
                message=SimpleNamespace(content=content), finish_reason='stop')])
        finally:
            with self.lock:
                self.active -= 1


class OCRPerformanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def images(self, count=6, duplicate=False):
        folder = self.root / 'images'
        folder.mkdir(exist_ok=True)
        for n in range(count):
            Image.new('RGB', (100, 80), (10 if duplicate else n*30, 20, 60)).save(folder / f'{n}.png')
        return folder

    def pdf(self, count=6):
        pages = [Image.new('RGB', (100, 80), (n*30, 20, 60)) for n in range(count)]
        path = self.root / 'scans.pdf'
        pages[0].save(path, save_all=True, append_images=pages[1:])
        return path

    def run_conversion(self, source, endpoint, **kwargs):
        with patch('tomarkdown.core.OpenAI', return_value=endpoint):
            return convert(source, api_key='test-only', base_url='http://unused.invalid/v1',
                           model='test-vision', results_dir=self.root/'out', **kwargs)

    def test_folder_requests_overlap_and_preserve_file_order(self):
        endpoint = VisionEndpoint()
        result = self.run_conversion(self.images(), endpoint)
        self.assertFalse(result.errors)
        self.assertEqual([p.stem for p in result.files], list(map(str, range(6))))
        self.assertGreater(endpoint.peak, 1, 'folder OCR still runs serially')
        self.assertLessEqual(endpoint.peak, 4)

    def test_pdf_requests_overlap_and_preserve_page_order(self):
        endpoint = VisionEndpoint()
        result = self.run_conversion(self.pdf(), endpoint)
        self.assertFalse(result.errors)
        text = result.files[0].read_text()
        self.assertEqual(text.count('OCR-'), 6)
        self.assertEqual(sorted(text.index(f'## Page {n}') for n in range(1,7)),
                         [text.index(f'## Page {n}') for n in range(1,7)])
        self.assertGreater(endpoint.peak, 1, 'PDF OCR still runs serially')
        self.assertLessEqual(endpoint.peak, 4)

    def test_duplicate_images_only_call_model_once(self):
        endpoint = VisionEndpoint()
        result = self.run_conversion(self.images(6, duplicate=True), endpoint)
        self.assertEqual(len(result.files), 6)
        self.assertEqual(endpoint.calls, 1)

    def test_pdf_ocr_failure_is_not_success(self):
        endpoint = VisionEndpoint(fail=True)
        result = self.run_conversion(self.pdf(1), endpoint)
        self.assertFalse(result.files, 'OCR failure was silently treated as success')
        self.assertEqual(len(result.errors), 1)
        self.assertFalse((self.root/'out'/'scans.md').exists())


    def test_mixed_files_share_global_budget_and_pdf_output_matches_baseline(self):
        source = self.images()
        self.pdf(6).rename(source / 'scans.pdf')
        endpoint = VisionEndpoint()
        result = self.run_conversion(source, endpoint, workers=3)
        self.assertFalse(result.errors)
        self.assertEqual(len(result.files), 7)
        self.assertGreater(endpoint.peak, 1)
        self.assertLessEqual(endpoint.peak, 3)
        baseline_endpoint = VisionEndpoint(delay=0)
        baseline = MarkItDown(enable_plugins=True, llm_client=baseline_endpoint,
                              llm_model='test-vision', llm_prompt=OCR_PROMPT,
                              ).convert(str(source/'scans.pdf')).text_content
        actual = (self.root/'out'/'images'/'scans.md').read_text().strip()
        self.assertEqual(actual, baseline.strip())

    def test_serial_and_concurrent_pdf_outputs_match_despite_reordered_completion(self):
        class Reordered(VisionEndpoint):
            def create(self, **kwargs):
                data = kwargs['messages'][0]['content'][1]['image_url']['url']
                time.sleep((int(hashlib.sha256(data.encode()).hexdigest()[:2],16)%4)*0.02)
                return super().create(**kwargs)
        path = self.pdf()
        serial = self.run_conversion(path, VisionEndpoint(delay=0), workers=1)
        text = serial.files[0].read_text()
        concurrent = self.run_conversion(path, Reordered(delay=0), workers=4)
        self.assertEqual(concurrent.files[0].read_text(), text)

    def test_pdf_is_opened_once(self):
        with patch('tomarkdown.pdf.pdfplumber.open', wraps=pdfplumber.open) as opened:
            result = self.run_conversion(self.pdf(), VisionEndpoint(delay=0))
        self.assertFalse(result.errors)
        self.assertEqual(opened.call_count, 1)

    def test_original_image_bytes_are_preserved(self):
        path = self.root/'large.png'
        Image.new('RGB', (800, 400), 'white').save(path)
        original = VisionEndpoint(delay=0)
        self.run_conversion(path, original)
        encoded = original.payloads[0]['messages'][0]['content'][1]['image_url']['url'].split(',',1)[1]
        self.assertEqual(base64.b64decode(encoded), path.read_bytes())
    def test_failure_is_not_cached(self):
        endpoint = VisionEndpoint(delay=0, fail=True)
        client = OCRClient(endpoint, 2)
        request = dict(model='test', messages=[{'content':[{'type':'text','text':'ocr'},
            {'type':'image_url','image_url':{'url':'data:image/png;base64,dGVzdA=='}}]}])
        with self.assertRaises(RuntimeError):
            client.create(**request)
        endpoint.fail = False
        client.reset_errors()
        client.create(**request)
        client.raise_if_failed()
        self.assertEqual(endpoint.calls, 2)

    def test_refusal_and_truncation_fail_but_blank_is_valid(self):
        for reason, content, refusal, should_fail in [
            ('length','partial',None,True), ('content_filter','',None,True),
            ('stop','', 'refused',True), ('stop',None,None,True), ('stop','',None,False),
        ]:
            endpoint = VisionEndpoint(delay=0)
            endpoint.chat.completions.create = lambda **kwargs: SimpleNamespace(choices=[SimpleNamespace(
                finish_reason=reason, message=SimpleNamespace(content=content,refusal=refusal))])
            client = OCRClient(endpoint, 1)
            with self.subTest(reason=reason,content=content,refusal=refusal):
                if should_fail:
                    with self.assertRaises(RuntimeError):
                        client.create(model='test',messages=[])
                else:
                    client.create(model='test',messages=[])
                    client.raise_if_failed()

    def test_same_stem_keeps_sorted_overwrite_behavior_and_no_ocr_never_calls_model(self):
        folder = self.root/'texts'
        folder.mkdir()
        (folder/'same.html').write_text('<p>first</p>')
        (folder/'same.txt').write_text('last')
        endpoint = VisionEndpoint()
        result = self.run_conversion(folder, endpoint, ocr=False)
        self.assertFalse(result.errors)
        self.assertEqual(endpoint.calls, 0)
        self.assertEqual(result.files[-1].read_text().strip(), 'last')

    def test_bad_performance_settings_are_rejected(self):
        for kwargs in [{'workers':0},{'workers':True},{'workers':33},
                       {'timeout':True},{'timeout':10**400},{'timeout':float('nan')},{'timeout':0},{'max_retries':-1}]:
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.run_conversion(self.root/'irrelevant',VisionEndpoint(),**kwargs)



    def test_text_pdf_does_not_call_ocr(self):
        import pymupdf
        path = self.root/'text.pdf'
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72,72), 'Preserve all native text 123.45')
            pdf.save(path)
        endpoint = VisionEndpoint(delay=0)
        result = self.run_conversion(path, endpoint)
        self.assertFalse(result.errors)
        self.assertEqual(endpoint.calls, 0)
        self.assertIn('Preserve all native text 123.45',result.files[0].read_text())

    def test_failed_embedded_image_extraction_uses_full_page_ocr(self):
        endpoint = VisionEndpoint(delay=0)
        with patch('tomarkdown.pdf._extract_images_from_page',return_value=[]):
            result = self.run_conversion(self.pdf(1),endpoint)
        self.assertFalse(result.errors)
        self.assertEqual(endpoint.calls,1)
        self.assertIn('OCR-',result.files[0].read_text())
        url=endpoint.payloads[0]['messages'][0]['content'][1]['image_url']['url']
        with Image.open(io.BytesIO(base64.b64decode(url.split(',',1)[1]))) as rendered:
            self.assertGreaterEqual(rendered.width,400)  # 100-point page at 300 DPI

    def test_office_ocr_failure_is_not_hidden_by_plugin(self):
        from docx import Document
        image=self.images(1)/'0.png'
        document=Document()
        document.add_paragraph('Native text')
        document.add_picture(str(image))
        source=self.root/'document.docx'
        document.save(source)
        result=self.run_conversion(source,VisionEndpoint(delay=0,fail=True))
        self.assertFalse(result.files)
        self.assertEqual(len(result.errors),1)


if __name__ == '__main__':
    unittest.main()

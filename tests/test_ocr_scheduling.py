"""A slow first request must not leave free workers idle behind it."""
from concurrent.futures import ThreadPoolExecutor
import threading
import unittest

import test_ocr_performance as fixtures


class GatedEndpoint(fixtures.VisionEndpoint):
    def __init__(self):
        super().__init__(delay=0)
        self.started = 0
        self.first_started = threading.Event()
        self.third_started = threading.Event()
        self.release_first = threading.Event()

    def create(self, **kwargs):
        with self.lock:
            self.started += 1
            order = self.started
        if order == 1:
            self.first_started.set()
            if not self.release_first.wait(5):
                raise RuntimeError('test gate timed out')
        if order == 3:
            self.third_started.set()
        return super().create(**kwargs)


class SchedulingTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.OCRPerformanceTests()
        self.fixture.setUp()

    def tearDown(self):
        self.fixture.tearDown()

    def assert_refills_before_first_finishes(self, source):
        endpoint = GatedEndpoint()
        with ThreadPoolExecutor(max_workers=1) as pool:
            result_future = pool.submit(self.fixture.run_conversion, source, endpoint, workers=2)
            try:
                self.assertTrue(endpoint.first_started.wait(2))
                refilled = endpoint.third_started.wait(0.5)
            finally:
                endpoint.release_first.set()
            result = result_future.result(timeout=5)
        self.assertFalse(result.errors)
        self.assertTrue(refilled, 'free worker stayed idle waiting for the first request')
        return result

    def test_pdf_refills_free_slot_before_slow_first_image_finishes(self):
        result = self.assert_refills_before_first_finishes(self.fixture.pdf(3))
        text = result.files[0].read_text()
        self.assertEqual(text.count('OCR-'), 3)
        self.assertLess(text.index('## Page 1'), text.index('## Page 2'))
        self.assertLess(text.index('## Page 2'), text.index('## Page 3'))

    def test_folder_refills_free_slot_before_slow_first_file_finishes(self):
        result = self.assert_refills_before_first_finishes(self.fixture.images(3))
        self.assertEqual([p.stem for p in result.files], ['0', '1', '2'])


    def test_docx_uses_available_ocr_workers(self):
        from docx import Document
        directory = self.fixture.images(3)
        document = Document()
        for image in sorted(directory.iterdir()):
            document.add_paragraph(image.stem)
            document.add_picture(str(image))
        source = self.fixture.root / 'images.docx'
        document.save(source)
        result = self.assert_refills_before_first_finishes(source)
        self.assertEqual(result.files[0].read_text().count('OCR-'), 3)


    def test_docx_output_and_image_payloads_match_original_plugin(self):
        from collections import Counter
        from docx import Document
        from markitdown import MarkItDown
        from tomarkdown.core import OCR_PROMPT
        directory = self.fixture.images(6)
        document = Document()
        for image in sorted(directory.iterdir()):
            document.add_paragraph('Paragraph ' + image.stem)
            document.add_picture(str(image))
        source = self.fixture.root / 'ordered.docx'
        document.save(source)
        baseline_endpoint = fixtures.VisionEndpoint(delay=0)
        expected = MarkItDown(enable_plugins=True, llm_client=baseline_endpoint,
                              llm_model='test-vision', llm_prompt=OCR_PROMPT).convert(str(source)).text_content
        endpoint = GatedEndpoint()
        with ThreadPoolExecutor(max_workers=1) as pool:
            result_future = pool.submit(self.fixture.run_conversion, source, endpoint, workers=3)
            try:
                self.assertTrue(endpoint.third_started.wait(2))
            finally:
                endpoint.release_first.set()
            result = result_future.result(timeout=5)
        self.assertFalse(result.errors)
        self.assertEqual(result.files[0].read_text(), expected + '\n')
        self.assertEqual(Counter(map(str, baseline_endpoint.payloads)), Counter(map(str, endpoint.payloads)))

    def test_file_reorder_buffer_remains_bounded(self):
        from tomarkdown.ocr import ordered_window
        release = threading.Event()
        fourth = threading.Event()
        fifth = threading.Event()
        def work(index):
            if index == 0:
                release.wait(3)
            if index == 3:
                fourth.set()
            if index == 4:
                fifth.set()
            return index
        with ThreadPoolExecutor(max_workers=2) as workers, ThreadPoolExecutor(max_workers=1) as caller:
            outcome = caller.submit(lambda: list(ordered_window(workers, work, range(20), 2)))
            try:
                self.assertTrue(fourth.wait(2))
                self.assertFalse(fifth.wait(0.1), 'completed results grew beyond the bounded buffer')
            finally:
                release.set()
            self.assertEqual(outcome.result(timeout=3), list(range(20)))


if __name__ == '__main__':
    unittest.main()
